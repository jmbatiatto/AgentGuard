import os
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, color_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def add_markdown_to_docx(doc, md_text):
    lines = md_text.split('\n')
    i = 0
    in_code_block = False
    code_lines = []
    
    while i < len(lines):
        line = lines[i]
        
        # Code block handling
        if line.strip().startswith('```'):
            if in_code_block:
                in_code_block = False
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Inches(0.4)
                p.paragraph_format.space_before = Pt(4)
                p.paragraph_format.space_after = Pt(4)
                run = p.add_run('\n'.join(code_lines))
                run.font.name = 'Consolas'
                run.font.size = Pt(9.5)
                run.font.color.rgb = RGBColor(30, 41, 59)
                code_lines = []
            else:
                in_code_block = True
                code_lines = []
            i += 1
            continue
            
        if in_code_block:
            code_lines.append(line)
            i += 1
            continue
            
        # Skip empty lines
        stripped = line.strip()
        if not stripped:
            i += 1
            continue
            
        # Horizontal rule
        if stripped in ('---', '***', '___'):
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(6)
            i += 1
            continue
            
        # Headings
        if stripped.startswith('# '):
            h = doc.add_heading(level=1)
            h.paragraph_format.space_before = Pt(14)
            h.paragraph_format.space_after = Pt(6)
            run = h.add_run(stripped[2:].strip())
            run.font.name = 'Segoe UI'
            run.font.size = Pt(20)
            run.font.bold = True
            run.font.color.rgb = RGBColor(15, 23, 42)
            i += 1
            continue
        elif stripped.startswith('## '):
            h = doc.add_heading(level=2)
            h.paragraph_format.space_before = Pt(12)
            h.paragraph_format.space_after = Pt(4)
            run = h.add_run(stripped[3:].strip())
            run.font.name = 'Segoe UI'
            run.font.size = Pt(15)
            run.font.bold = True
            run.font.color.rgb = RGBColor(30, 58, 138)
            i += 1
            continue
        elif stripped.startswith('### '):
            h = doc.add_heading(level=3)
            h.paragraph_format.space_before = Pt(10)
            h.paragraph_format.space_after = Pt(3)
            run = h.add_run(stripped[4:].strip())
            run.font.name = 'Segoe UI'
            run.font.size = Pt(12.5)
            run.font.bold = True
            run.font.color.rgb = RGBColor(71, 85, 105)
            i += 1
            continue
        elif stripped.startswith('#### '):
            h = doc.add_heading(level=4)
            h.paragraph_format.space_before = Pt(8)
            h.paragraph_format.space_after = Pt(2)
            run = h.add_run(stripped[5:].strip())
            run.font.name = 'Segoe UI'
            run.font.size = Pt(11)
            run.font.bold = True
            run.font.color.rgb = RGBColor(100, 116, 139)
            i += 1
            continue

        # Markdown Table
        if '|' in stripped and i + 1 < len(lines) and '|' in lines[i+1] and ('---' in lines[i+1] or ':---' in lines[i+1]):
            table_lines = []
            while i < len(lines) and '|' in lines[i].strip():
                table_lines.append(lines[i].strip())
                i += 1
            
            # Parse table
            if len(table_lines) >= 3:
                headers = [c.strip() for c in table_lines[0].strip('|').split('|')]
                data_rows = []
                for tline in table_lines[2:]:
                    cols = [c.strip() for c in tline.strip('|').split('|')]
                    data_rows.append(cols)
                
                table = doc.add_table(rows=len(data_rows) + 1, cols=len(headers))
                table.alignment = WD_TABLE_ALIGNMENT.CENTER
                
                # Header row
                hdr_cells = table.rows[0].cells
                for idx, text in enumerate(headers):
                    if idx < len(hdr_cells):
                        cell = hdr_cells[idx]
                        set_cell_background(cell, "1E3A8A")
                        set_cell_margins(cell, top=120, bottom=120, left=140, right=140)
                        p = cell.paragraphs[0]
                        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                        run = p.add_run(text)
                        run.font.bold = True
                        run.font.name = 'Segoe UI'
                        run.font.size = Pt(9.5)
                        run.font.color.rgb = RGBColor(255, 255, 255)
                
                # Data rows
                for r_idx, row_data in enumerate(data_rows):
                    row_cells = table.rows[r_idx + 1].cells
                    bg_color = "F8FAFC" if r_idx % 2 == 1 else "FFFFFF"
                    for c_idx, text in enumerate(row_data):
                        if c_idx < len(row_cells):
                            cell = row_cells[c_idx]
                            set_cell_background(cell, bg_color)
                            set_cell_margins(cell, top=80, bottom=80, left=140, right=140)
                            p = cell.paragraphs[0]
                            # Simple formatting for text
                            clean_text = re.sub(r'\*\*(.*?)\*\*', r'\1', text)
                            clean_text = re.sub(r'`(.*?)`', r'\1', clean_text)
                            run = p.add_run(clean_text)
                            run.font.name = 'Segoe UI'
                            run.font.size = Pt(9)
                            run.font.color.rgb = RGBColor(30, 41, 59)
                
                doc.add_paragraph().paragraph_format.space_after = Pt(6)
            continue

        # Blockquote
        if stripped.startswith('> '):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.4)
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(4)
            quote_text = stripped[2:].strip()
            # Clean formatting markers
            quote_text = quote_text.replace('> ', ' ')
            quote_text = re.sub(r'\[!(NOTE|TIP|IMPORTANT|WARNING|CAUTION)\]', r'[\1]', quote_text)
            run = p.add_run(quote_text)
            run.font.italic = True
            run.font.name = 'Segoe UI'
            run.font.size = Pt(10)
            run.font.color.rgb = RGBColor(71, 85, 105)
            i += 1
            continue

        # Bullet list item
        if stripped.startswith('- ') or stripped.startswith('* ') or re.match(r'^\d+\.\s', stripped):
            p = doc.add_paragraph(style='List Bullet' if (stripped.startswith('- ') or stripped.startswith('* ')) else 'List Number')
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            
            if stripped.startswith('- ') or stripped.startswith('* '):
                item_text = stripped[2:].strip()
            else:
                item_text = re.sub(r'^\d+\.\s', '', stripped).strip()
            
            # Format bold within paragraph
            parts = re.split(r'(\*\*.*?\*\*)', item_text)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    run = p.add_run(part[2:-2])
                    run.font.bold = True
                else:
                    clean_part = re.sub(r'`(.*?)`', r'\1', part)
                    run = p.add_run(clean_part)
                run.font.name = 'Segoe UI'
                run.font.size = Pt(10)
                run.font.color.rgb = RGBColor(30, 41, 59)
            i += 1
            continue

        # Normal paragraph
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        
        parts = re.split(r'(\*\*.*?\*\*)', stripped)
        for part in parts:
            if part.startswith('**') and part.endswith('**'):
                run = p.add_run(part[2:-2])
                run.font.bold = True
            else:
                clean_part = re.sub(r'`(.*?)`', r'\1', part)
                run = p.add_run(clean_part)
            run.font.name = 'Segoe UI'
            run.font.size = Pt(10)
            run.font.color.rgb = RGBColor(30, 41, 59)
        i += 1

def build_docx_files():
    base_dir = r"c:\Users\jonat\Downloads\Codigos\AgentGuard\01-product"
    files = [
        ("product-brief.md", "product-brief.docx", "AgentGuard — Product Brief v2.0"),
        ("requirements.md", "requirements.docx", "AgentGuard — Requisitos de Software (SRS) v2.0"),
        ("user-stories.md", "user-stories.docx", "AgentGuard — Historias de Usuario (Agile Backlog) v2.0"),
        ("use-cases.md", "use-cases.docx", "AgentGuard — Casos de Uso del Sistema v2.0")
    ]
    
    # 1. Generate individual docx files
    for md_name, docx_name, doc_title in files:
        md_path = os.path.join(base_dir, md_name)
        docx_path = os.path.join(base_dir, docx_name)
        
        if os.path.exists(md_path):
            with open(md_path, 'r', encoding='utf-8') as f:
                content = f.read()
            
            doc = Document()
            # Set margins to normal 1 inch
            for section in doc.sections:
                section.top_margin = Inches(1)
                section.bottom_margin = Inches(1)
                section.left_margin = Inches(1)
                section.right_margin = Inches(1)
            
            add_markdown_to_docx(doc, content)
            doc.save(docx_path)
            print(f"Generated: {docx_name}")

    # 2. Generate Master Document (all 4 in one consolidated file for Google Docs upload)
    master_doc = Document()
    for section in master_doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
        
    for idx, (md_name, _, _) in enumerate(files):
        md_path = os.path.join(base_dir, md_name)
        if os.path.exists(md_path):
            with open(md_path, 'r', encoding='utf-8') as f:
                content = f.read()
            if idx > 0:
                master_doc.add_page_break()
            add_markdown_to_docx(master_doc, content)
            
    master_path = os.path.join(base_dir, "AgentGuard_Documentacion_Producto_Completa.docx")
    master_doc.save(master_path)
    print("Generated Master Document: AgentGuard_Documentacion_Producto_Completa.docx")

if __name__ == "__main__":
    build_docx_files()
